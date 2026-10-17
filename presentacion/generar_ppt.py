"""Generador de la presentacion de la defensa EP2 (PPT).

Uso (con python-pptx instalado):
    pip install python-pptx
    python presentacion/generar_ppt.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = Path(__file__).resolve().parent / "Presentacion_EP2_Ferrandes.pptx"

AZUL = RGBColor(0x1F, 0x4E, 0x79)
CELESTE = RGBColor(0x2E, 0x86, 0xDE)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
GRIS = RGBColor(0x60, 0x60, 0x60)
NARANJA = RGBColor(0xE6, 0x7E, 0x22)


def nueva_presentacion() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def diapositiva(prs, titulo: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # en blanco
    # franja superior
    superior = slide.shapes.add_textbox(0, 0, prs.slide_width, Inches(1.0))
    tf = superior.text_frame
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = titulo
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = BLANCO
    # fondo de franja via rectangulo
    fondo = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.0))  # MSO_SHAPE.RECTANGLE
    fondo.fill.solid()
    fondo.fill.fore_color.rgb = AZUL
    fondo.line.fill.background()
    fondo.shadow.inherit = False
    # reordenar: poner el fondo detras (el texto se dibuja despues de nuevo)
    slide.shapes._spTree.remove(fondo._element)
    slide.shapes._spTree.insert(2, fondo._element)
    return slide


def caja_bullets(slide, items, left, top, width, height, tam=16, color=GRIS):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    primera = True
    for item in items:
        p = tf.paragraphs[0] if primera else tf.add_paragraph()
        primera = False
        p.text = item
        p.font.size = Pt(tam)
        p.font.color.rgb = color
        p.space_after = Pt(8)
    return tb


def generar():
    prs = nueva_presentacion()

    # ------------------------------------------------- 1 portada
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fondo = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    fondo.fill.solid(); fondo.fill.fore_color.rgb = AZUL; fondo.line.fill.background(); fondo.shadow.inherit = False
    tb = slide.shapes.add_textbox(Inches(1), Inches(1.6), Inches(11), Inches(4))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Agente Ferrandes"
    p.font.size = Pt(54); p.font.bold = True; p.font.color.rgb = BLANCO
    p = tf.add_paragraph(); p.text = "Automatización del flujo B2B — Ferretería Industrial Andes SpA"
    p.font.size = Pt(22); p.font.color.rgb = CELESTE
    p = tf.add_paragraph(); p.text = "Evaluación Parcial N°2 · ISY0101 · Ingeniería de Soluciones con IA"
    p.font.size = Pt(18); p.font.color.rgb = BLANCO
    p = tf.add_paragraph(); p.text = "Integrantes: ⬅ completar  |  Docente: Juan Burgueño"
    p.font.size = Pt(16); p.font.color.rgb = CELESTE

    # ------------------------------------------------- 2 contexto
    s = diapositiva(prs, "Contexto y problema")
    caja_bullets(s, [
        "FERRANDES (Concepción): venta y distribución B2B de tornillería, herramientas, EPP y materiales.",
        "Proceso actual manual: stock → descuento/IVA → cotización → crédito → confirmación.",
        "Demoras de horas, errores de cálculo y sin trazabilidad (SLA 48h de despacho).",
        "EP1: prototipo RAG de asistente (respuestas).",
        "EP2: agente FUNCIONAL que ejecuta acciones y decide (LangChain).",
    ], 0.7, 1.3, 12, 5, tam=18)

    # ------------------------------------------------- 3 solución
    s = diapositiva(prs, "Solución: agente orquestador")
    caja_bullets(s, [
        "3 familias de herramientas: CONSULTA · ESCRITURA · RAZONAMIENTO (13 tools).",
        "Memoria de corto plazo (sesión) + largo plazo (SQLite) para flujos prolongados.",
        "Planificador: esquema priorizado de 6 etapas con inyección adaptativa.",
        "Framework: LangChain create_agent (grafo start→model→tools→model→end).",
        "LLM real (Claude) o modelo simulado sin costo para evaluar.",
    ], 0.7, 1.3, 12, 5, tam=18)

    # ------------------------------------------------- 4 arquitectura
    s = diapositiva(prs, "Arquitectura y orquestación")
    s.shapes.add_picture(str(RAIZ / "assets" / "diagrama_arquitectura.png"),
                         Inches(3.0), Inches(1.2), height=Inches(5.6))

    # ------------------------------------------------- 5 herramientas
    s = diapositiva(prs, "Herramientas del agente (IL2.1)")
    datos = [
        ("CONSULTA", "stock, precio, catálogo, ficha, pedido, cotización", "lectura sin efectos"),
        ("ESCRITURA", "cotización, pedido, orden de reposición, ticket", "persistente + trazabilidad"),
        ("RAZONAMIENTO", "riesgo de crédito, reposición, plan de pedido", "decisión adaptativa"),
    ]
    top = 1.4
    for nombre, lista, efecto in datos:
        tb = s.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12), Inches(1.5))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = nombre
        p.font.size = Pt(20); p.font.bold = True; p.font.color.rgb = NARANJA
        p = tf.add_paragraph(); p.text = f"{lista} → {efecto}"
        p.font.size = Pt(16); p.font.color.rgb = GRIS
        top += 1.85

    # ------------------------------------------------- 6 memoria
    s = diapositiva(prs, "Memoria y recuperación de contexto (IL2.2)")
    caja_bullets(s, [
        "Corto plazo: buffer por sesión + resumen comprimido (coherencia en el diálogo).",
        "Largo plazo (SQLite): clientes, cotizaciones, pedidos con historial de etapas, tickets y OC.",
        "Continuidad: cotización de la sesión A → confirmación/estado en la sesión B (persistencia real).",
        "Recuperación de contexto: retriever semántico TF-IDF + embeddings densos opcionales (IE4).",
        "Trazabilidad: evidencia/logs_agente.jsonl con fuentes y decisiones por consulta.",
    ], 0.7, 1.3, 12, 5, tam=17)

    # ------------------------------------------------- 6b metricas del retriever
    s = diapositiva(prs, "Recuperación semántica: mejora y métricas (IE4)")
    caja_bullets(s, [
        "Similitud coseno sobre vectores TF-IDF (pondera términos discriminantes y normaliza).",
        "Normalización morfológica ligera: “ubicados” empareja con “ubicación”.",
        "Embeddings densos opcionales (sentence-transformers) sin cambiar la interfaz.",
        "Medición objetiva: evaluar_recuperacion() → hit-rate@k y MRR.",
        "Resultado: hit-rate@3 = 100 % · MRR = 1,00 (8 consultas etiquetadas).",
    ], 0.7, 1.3, 12, 5, tam=17)

    # ------------------------------------------------- 7 planificación
    s = diapositiva(prs, "Planificación y decisiones adaptativas (IL2.3)")
    s.shapes.add_picture(str(RAIZ / "assets" / "diagrama_flujo_pedido.png"),
                         Inches(0.5), Inches(1.15), height=Inches(6.0))
    caja_bullets(s, [
        "Plan priorizado por criticidad (6 etapas).",
        "Stock bajo → OC automática + entrega programada.",
        "Monto > cupo/umbral → EN_APROBACION + ticket.",
        "Cliente nuevo → pago contra entrega.",
        "Normativa EPP → deriva a especialista (guardrail).",
        "Ítems inválidos → REPLANIFICA y continúa.",
    ], 4.6, 1.35, 8.2, 5.5, tam=15, color=NARANJA)

    # ------------------------------------------------- 7b justificacion
    s = diapositiva(prs, "Justificación de la orquestación (IE8, IE9)")
    caja_bullets(s, [
        "LangChain create_agent: grafo oficial, documentado, cambia de modelo (Claude o simulado).",
        "Herramientas puras por familia (consulta/escritura/razonamiento): separación de efectos y testeo.",
        "SQLite como memoria de largo plazo: persistencia sin servidor, migrable a PostgreSQL.",
        "Retriever TF-IDF: precisión medible hoy y ruta directa a embeddings/vector store.",
        "Planificador priorizado: cada paso ligado a un requisito del flujo (stock, crédito, SLA).",
    ], 0.7, 1.3, 12, 5, tam=16)

    # ------------------------------------------------- 8 demo
    s = diapositiva(prs, "Demo en vivo")
    caja_bullets(s, [
        "1) Consulta de stock y precio (herramienta de consulta).",
        "2) Pedido con faltante TOR-002: el agente emite OC y cotiza (adaptación).",
        "3) Pedido que supera cupo: escalamiento a vendedor con ticket.",
        "4) Ítems inválidos: replanifica y continúa con los válidos.",
        "Comando: python scripts/demo.py   |   python tests/ejecutar_tests.py",
        "Respaldo si falla en vivo: evidencia/ + notebook de Colab ya ejecutado.",
    ], 0.7, 1.25, 12, 5.2, tam=16)

    # ------------------------------------------------- 9 evidencias
    s = diapositiva(prs, "Evidencia y resultados")
    caja_bullets(s, [
        "38 pruebas unitarias validadas (tests/ejecutar_tests.py) — todas 'ok'.",
        "11 escenarios de negocio en evidencia/evidencia_pruebas.txt.",
        "Retriever: hit-rate@3 = 100 % · MRR = 1,00 (8 consultas etiquetadas).",
        "Flujos: PEDIDO_CONFIRMADO, EN_APROBACION, DERIVADO, REPLANIFICADO.",
        "README + diagramas de orquestación + informe técnico (≤5 páginas).",
    ], 0.7, 1.3, 12, 5, tam=17)

    # ------------------------------------------------- 9b retos y limites
    s = diapositiva(prs, "Retos, límites y próximos pasos")
    caja_bullets(s, [
        "Límite: el grafo LangChain se validó con un LLM simulado (sin API key en la sala).",
        "Límite: el retriever es TF-IDF; los embeddings densos son opcionales.",
        "Próximo paso: embeddings reales + vector store (Chroma/FAISS) y medición de calidad.",
        "Próximo paso: memoria LangChain explícita en el grafo (checkpointer) y traza con LLM real.",
        "Próximo paso: despliegue con PostgreSQL y pruebas con datos de producción.",
    ], 0.7, 1.3, 12, 5, tam=17)

    # ------------------------------------------------- 10 cierre
    s = diapositiva(prs, "Cierre")
    caja_bullets(s, [
        "Decisiones de diseño justificadas por alineación con el flujo (IE9).",
        "Autonomía demostrada: decisión adaptativa ante stock, crédito, normativa e ítems inválidos.",
        "Repositorio: código + documentación + evidencia reproducible (38 pruebas, hit-rate@3 100 %).",
        "Siguiente paso: embeddings reales (Chroma/FAISS), memoria LangChain y despliegue.",
        "¡Gracias! Preguntas →",
    ], 0.7, 1.3, 12, 5, tam=18)

    prs.save(SALIDA)
    print("📊 Presentación generada:", SALIDA)


if __name__ == "__main__":
    generar()